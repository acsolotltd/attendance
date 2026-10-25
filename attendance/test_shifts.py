"""Django-free tests of the shift engine: python -m unittest attendance.test_shifts"""
import unittest
from datetime import date, datetime, time, timedelta

from attendance import shifts

DAY = {"id": 1, "name": "Day", "start": time(8, 0), "end": time(17, 0), "days": {0, 1, 2, 3, 4},
       "grace": 10, "ot_after": 30, "before": 120, "after": 180, "is_default": True}
NIGHT = dict(DAY, id=2, name="Night", start=time(20, 0), end=time(6, 0), days={0, 1, 2, 3, 4, 5}, is_default=False)
D = date(2026, 10, 5)  # a Monday


def punch(uid, ts, name=""):
    return {"user_id": uid, "name": name, "timestamp": ts}


def resolver(mapping, default=None):
    return lambda uid, day: mapping.get(uid, default)


class ShiftTests(unittest.TestCase):
    def run_report(self, records, mapping, default=None, start=D, end=D, **kw):
        return shifts.build_shift_report(records, resolver(mapping, default), start, end,
                                         now=kw.pop("now", datetime(2026, 12, 1)), **kw)

    def test_day_shift_on_time_with_overtime(self):
        rows, _ = self.run_report([punch("1", datetime(2026, 10, 5, 7, 55)), punch("1", datetime(2026, 10, 5, 17, 40))],
                                  {"1": DAY})
        self.assertEqual(len(rows), 1)
        r = rows[0]
        self.assertEqual((r["late_min"], r["ot_min"], r["hours"]), (0, 40, 9.75))
        self.assertEqual(r["status"], "Overtime")

    def test_night_shift_crossing_midnight_is_one_row_dated_by_start(self):
        recs = [punch("2", datetime(2026, 10, 5, 19, 58)), punch("2", datetime(2026, 10, 6, 6, 7))]
        rows, stats = self.run_report(recs, {"2": NIGHT})
        self.assertEqual([(r["date"], r["status"]) for r in rows if r["status"] != "Absent"], [(D, "On time")])
        self.assertEqual(rows[0]["hours"], 10.15)
        self.assertEqual(stats["unmatched_punches"], 0)

    def test_late_and_early_leave(self):
        recs = [punch("2", datetime(2026, 10, 5, 20, 30)), punch("2", datetime(2026, 10, 6, 5, 0))]
        r = self.run_report(recs, {"2": NIGHT})[0][0]
        self.assertEqual((r["late_min"], r["early_min"]), (30, 60))
        self.assertEqual(r["status"], "Late, Early leave")

    def test_single_punch_is_missing_checkout_or_checkin(self):
        r = self.run_report([punch("1", datetime(2026, 10, 5, 8, 2))], {"1": DAY})[0][0]
        self.assertEqual((r["status"], r["out"]), ("Missing check-out", None))
        r = self.run_report([punch("1", datetime(2026, 10, 5, 17, 2))], {"1": DAY})[0][0]
        self.assertEqual((r["status"], r["in"]), ("Missing check-in", None))

    def test_double_tap_counts_once(self):
        recs = [punch("1", datetime(2026, 10, 5, 7, 58)), punch("1", datetime(2026, 10, 5, 7, 59))]
        self.assertEqual(self.run_report(recs, {"1": DAY})[0][0]["status"], "Missing check-out")

    def test_shift_date_outside_range_is_excluded(self):
        recs = [punch("2", datetime(2026, 10, 4, 20, 0)), punch("2", datetime(2026, 10, 5, 6, 0))]  # started Oct 4
        rows, _ = self.run_report(recs, {"2": NIGHT}, start=D, end=D)
        self.assertEqual([r for r in rows if r["status"] != "Absent"], [])

    def test_absent_only_on_past_workdays(self):
        rows, _ = self.run_report([], {"1": DAY}, start=date(2026, 10, 3), end=date(2026, 10, 6),
                                  extra_users={"1"}, now=datetime(2026, 10, 6, 12, 0))
        self.assertEqual([(r["date"], r["status"]) for r in rows],
                         [(date(2026, 10, 5), "Absent"), (date(2026, 10, 6), "Absent")])  # Sat/Sun are rest days
        future, _ = self.run_report([], {"1": DAY}, extra_users={"1"}, now=datetime(2026, 10, 5, 6, 0))
        self.assertEqual(future, [])  # shift hasn't started yet

    def test_user_without_shift_reported(self):
        rows, stats = self.run_report([punch("9", datetime(2026, 10, 5, 8, 0))], {})
        self.assertEqual((rows, stats["users_without_shift"]), ([], ["9"]))

    def test_default_shift_applies_to_everyone(self):
        recs = [punch("7", datetime(2026, 10, 5, 8, 0), "Ann"), punch("7", datetime(2026, 10, 5, 17, 0))]
        rows, _ = self.run_report(recs, {}, default=DAY)
        self.assertEqual((rows[0]["name"], rows[0]["shift"]), ("Ann", "Day"))

    def test_rotation_blocks_and_totals(self):
        blocks = shifts.rotation_blocks(date(2026, 10, 5), date(2026, 10, 25), 1)
        self.assertEqual([(b[0], b[1], b[2]) for b in blocks],
                         [(date(2026, 10, 5), date(2026, 10, 11), 0), (date(2026, 10, 12), date(2026, 10, 18), 1),
                          (date(2026, 10, 19), date(2026, 10, 25), 2)])
        recs = [punch("1", datetime(2026, 10, 5, 8, 30)), punch("1", datetime(2026, 10, 5, 17, 0))]
        rows, _ = self.run_report(recs, {"1": DAY}, extra_users={"1"}, end=date(2026, 10, 6))
        t = shifts.totals(rows)[0]
        self.assertEqual((t["shifts"], t["absent"], t["late_count"], t["late_min"]), (1, 1, 1, 30))


if __name__ == "__main__":
    unittest.main()


class WorkbookTests(unittest.TestCase):
    def test_workbook_has_shift_sheets_with_night_shift_dates(self):
        import io

        import openpyxl

        from attendance import services
        recs = [punch("2", datetime(2026, 10, 5, 19, 58), "Ann"), punch("2", datetime(2026, 10, 6, 6, 7))]
        rows, stats = shifts.build_shift_report(recs, resolver({"2": NIGHT}), D, D, now=datetime(2026, 12, 1))
        buf = services.build_workbook([], [], D, D, shift_rows=rows, shift_stats=stats)
        wb = openpyxl.load_workbook(io.BytesIO(buf.getvalue()))
        self.assertIn("Shift Report", wb.sheetnames)
        self.assertIn("Shift Totals", wb.sheetnames)
        data = list(wb["Shift Report"].iter_rows(min_row=2, values_only=True))
        self.assertEqual((data[0][1], data[0][3], data[0][12]), ("2", "Night", "On time"))
        self.assertEqual(data[0][7], datetime(2026, 10, 6, 6, 7))  # check-out is on the next calendar day
        info = dict(wb["Report Info"].iter_rows(min_row=2, values_only=True))
        self.assertEqual(info["Users with no shift"], "none")
