"""ZKTeco / MB360 access layer (pyzk over TCP 4370) plus Excel/CSV builders.

Deliberately Django-free so it can be unit-tested without a database.
"""
import csv
import io
import logging
import random
import re
import socket
import threading
import time
from collections import defaultdict
from concurrent.futures import ThreadPoolExecutor
from contextlib import contextmanager
from datetime import datetime, timedelta
from types import SimpleNamespace

from . import shifts as shiftlib
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

logger = logging.getLogger(__name__)

RETRIES = 2
PUNCH_LABELS = {0: "Check In", 1: "Check Out", 2: "Break Out", 3: "Break In", 4: "Overtime In", 5: "Overtime Out"}
VERIFY_LABELS = {0: "Password", 1: "Fingerprint", 2: "Card", 3: "Password", 4: "Card", 5: "Fingerprint+Card",
                 6: "Fingerprint+Password", 7: "Card+Password", 8: "Card+Fingerprint+Password", 15: "Face", 25: "Palm"}
ADMS_ONLINE_SECONDS = 150  # ADMS devices poll roughly every 30 s; offline after ~5 missed polls
DEMO_NAMES = ["Sade Williams", "Chidi Okafor", "Tunde Adeyemi", "Ngozi Eze", "Ibrahim Musa",
              "Funke Adebayo", "Emeka Nwosu", "Zainab Yusuf", "Segun Balogun", "Kemi Ogunleye"]


def label(mapping, code):
    return mapping.get(code, f"Code {code}")


def is_demo(device):
    return str(device.ip_address).strip().lower().startswith("demo")


def friendly_error(exc):
    msg = f"{type(exc).__name__}: {exc}".strip().rstrip(":")
    low = msg.lower()
    if "uthenticat" in low:
        hint = "wrong Comm Key?"
    elif "timed out" in low or "timeout" in low:
        hint = "no/slow reply — check port forwarding, or raise this device's timeout"
    elif "refused" in low:
        hint = "connection refused — wrong port or forwarding rule?"
    elif "can't connect" in low or "cant connect" in low:
        hint = "handshake rejected — wrong Comm Key, or another program (e.g. ZKBioTime) is connected"
    elif "no module named" in low:
        hint = "run: pip install pyzk"
    else:
        hint = ""
    return f"{msg} — {hint}" if hint else msg


# --------------------------------------------------------------------------- connection

_LOCKS, _LOCKS_GUARD = {}, threading.Lock()


def _lock_for(device):
    """A device accepts one connection at a time: serialize our own sessions to the same ip:port."""
    key = (str(device.ip_address).strip().lower(), int(device.port))
    with _LOCKS_GUARD:
        return _LOCKS.setdefault(key, threading.Lock())


def _make_zk(device, force_udp=False, timeout=None):
    from zk import ZK  # imported lazily so a missing pyzk is reported per-device, not at start-up
    kwargs = dict(
        port=int(device.port),
        timeout=int(timeout or device.timeout or 15),
        password=int(device.comm_key or 0),
        force_udp=force_udp,
        ommit_ping=True,  # public IPs often block ICMP
    )
    host = str(device.ip_address).strip()
    try:
        return ZK(host, encoding=getattr(device, "encoding", "") or "UTF-8", **kwargs)
    except TypeError:  # older pyzk without the encoding argument
        return ZK(host, **kwargs)


def probe(device, timeout=3):
    """Returns (online, latency_ms). ADMS: heard from recently. TCP: port test. UDP: real handshake."""
    if getattr(device, "mode", "pull") == "adms":
        last = getattr(device, "last_seen", None)
        return bool(last) and (datetime.now() - last).total_seconds() <= ADMS_ONLINE_SECONDS, None
    if is_demo(device):
        return True, 0
    started = time.monotonic()
    if getattr(device, "protocol", "tcp") == "udp":
        try:
            conn = _make_zk(device, force_udp=True, timeout=4).connect()
            conn.disconnect()
            return True, int((time.monotonic() - started) * 1000)
        except Exception:
            return False, None
    try:
        with socket.create_connection((str(device.ip_address).strip(), int(device.port)), timeout=timeout):
            return True, int((time.monotonic() - started) * 1000)
    except (OSError, ValueError):
        return False, None


def _connect(device):
    use_udp = getattr(device, "protocol", "tcp") == "udp"
    if not use_udp and not probe(device)[0]:
        raise ConnectionError(
            f"Cannot reach {device.ip_address}:{device.port} over TCP — check port forwarding, "
            "firewall, device power, or try the UDP protocol"
        )
    last = None
    for attempt in range(RETRIES):
        try:
            return _make_zk(device, force_udp=use_udp).connect()
        except Exception as exc:
            last = exc
            logger.warning("connect attempt %s to %s failed: %s", attempt + 1, device.name, exc)
            if "uthenticat" in str(exc).lower():
                break  # retrying with the same Comm Key cannot help
            time.sleep(1.5)
    raise last


@contextmanager
def device_session(device, lock=True):
    """Connects, optionally pauses the device while we read, and always restores/disconnects."""
    with _lock_for(device):
        conn = _connect(device)
        try:
            if lock:
                try:
                    conn.disable_device()
                except Exception:
                    logger.warning("could not disable %s during download", device.name)
            yield conn
        finally:
            if lock:
                try:
                    conn.enable_device()
                except Exception:
                    pass
            try:
                conn.disconnect()
            except Exception:
                pass


# --------------------------------------------------------------------------- demo simulator

def _demo_data(device, days=60):
    """Deterministic fake workplace punches (incl. a few duplicates) for testing without hardware."""
    rng = random.Random(f"demo:{device.name}")
    users = [SimpleNamespace(user_id=str(101 + i), name=n) for i, n in enumerate(DEMO_NAMES)]
    now = datetime.now()
    logs = []
    for back in range(days, -1, -1):
        day = (now - timedelta(days=back)).replace(hour=0, minute=0, second=0, microsecond=0)
        if day.weekday() >= 5:
            continue
        for u in users:
            if rng.random() < 0.08:
                continue  # absent
            verify = rng.choice([1, 1, 15, 2])
            events = [(timedelta(hours=8, minutes=rng.randint(-40, 45)), 0)]
            if rng.random() < 0.8:
                events.append((timedelta(hours=12, minutes=rng.randint(20, 50)), 2))
                events.append((timedelta(hours=13, minutes=rng.randint(10, 50)), 3))
            if rng.random() < 0.93:
                events.append((timedelta(hours=17, minutes=rng.randint(-30, 90)), 1))
            for delta, punch in events:
                ts = day + delta + timedelta(seconds=rng.randint(0, 59))
                if ts <= now:
                    logs.append(SimpleNamespace(user_id=u.user_id, timestamp=ts, status=verify, punch=punch))
    if logs:
        logs.extend(rng.sample(logs, min(3, len(logs))))  # exact duplicates exercise de-duplication
    return users, logs


# --------------------------------------------------------------------------- download

def download_device_logs(device):
    """Raw download. Returns ({user_id: name}, [log objects])."""
    if is_demo(device):
        users, logs = _demo_data(device)
        return {u.user_id: u.name for u in users}, logs
    with device_session(device, lock=True) as conn:
        try:
            users = {str(u.user_id).strip(): (u.name or "").strip() for u in conn.get_users()}
        except Exception as exc:
            logger.warning("get_users failed on %s: %s", device.name, exc)
            users = {}
        logs = conn.get_attendance() or []
    return users, logs


def normalize_logs(device, users, logs):
    """Turns raw pyzk logs into dicts, removing exact duplicates. Returns (records, duplicates_removed)."""
    seen, out, dupes = set(), [], 0
    for log in logs:
        ts = getattr(log, "timestamp", None)
        if not isinstance(ts, datetime):
            continue
        uid = str(log.user_id).strip()
        if (uid, ts) in seen:
            dupes += 1
            continue
        seen.add((uid, ts))
        punch = int(getattr(log, "punch", 0) or 0)
        verify = int(getattr(log, "status", 0) or 0)
        out.append({
            "device": device.name, "user_id": uid, "name": users.get(uid, ""), "timestamp": ts,
            "punch": "Punch" if getattr(device, "punch_state_unused", False) else label(PUNCH_LABELS, punch),
            "punch_code": punch,
            "verify": label(VERIFY_LABELS, verify), "verify_code": verify,
        })
    out.sort(key=lambda r: (r["timestamp"], r["user_id"]))
    return out, dupes


def parse_user_filter(text):
    """'101, 105; John' -> ['101', '105', 'john'] (comma/semicolon/newline separated)."""
    return [t.strip().lower() for t in re.split(r"[,;\n]+", text or "") if t.strip()]


def matches_user(tokens, user_id, name):
    uid, nm = str(user_id).lower(), (name or "").lower()
    return any(t == uid or t in nm for t in tokens)


def filter_records(records, start, end, tokens=None):
    return [
        r for r in records
        if start <= r["timestamp"].date() <= end and (not tokens or matches_user(tokens, r["user_id"], r["name"]))
    ]


def fetch_device_records(device, start, end, tokens=None, keep_all=False):
    """Never raises. Returns (records_in_range, meta). meta['all_records'] holds the full log if keep_all."""
    started = time.monotonic()
    meta = {"name": device.name, "address": f"{device.ip_address}:{device.port}", "count": 0, "on_device": 0,
            "duplicates": 0, "error": None, "elapsed": 0.0, "source": "device"}
    records = []
    try:
        users, logs = download_device_logs(device)
        everything, dupes = normalize_logs(device, users, logs)
        records = filter_records(everything, start, end, tokens)
        meta.update(count=len(records), on_device=len(everything), duplicates=dupes)
        if keep_all:
            meta["all_records"] = everything
    except Exception as exc:
        logger.warning("download from %s failed: %s", device.name, exc)
        meta["error"] = friendly_error(exc)
    meta["elapsed"] = round(time.monotonic() - started, 1)
    return records, meta


def fetch_all(devices, start, end, tokens=None, keep_all=False):
    devices = list(devices)
    if not devices:
        return [], []
    with ThreadPoolExecutor(max_workers=min(8, len(devices))) as pool:
        results = list(pool.map(lambda d: fetch_device_records(d, start, end, tokens, keep_all), devices))
    records, summary = [], []
    for recs, meta in results:
        records.extend(recs)
        summary.append(meta)
    records.sort(key=lambda r: (r["timestamp"], r["device"], r["user_id"]))
    return records, summary


def _fmt_delta(seconds):
    sign = "+" if seconds >= 0 else "-"
    minutes = int(abs(seconds) // 60)
    return f"{sign}{minutes // 60}h {minutes % 60}m"


def device_info(device):
    """Full handshake + device details, for the 'Test' button. Never raises."""
    started = time.monotonic()
    if is_demo(device):
        _, logs = _demo_data(device)
        return {"ok": True, "elapsed": 0.0, "note": "", "rows": [
            ("Mode", "Built-in simulator (no hardware)"), ("Users", len(DEMO_NAMES)), ("Attendance records", len(logs))]}
    try:
        rows, note = [], ""
        with device_session(device, lock=False) as conn:
            for title, method in (("Device name", "get_device_name"), ("Serial number", "get_serialnumber"),
                                  ("Firmware", "get_firmware_version"), ("Platform", "get_platform"),
                                  ("MAC", "get_mac")):
                try:
                    value = getattr(conn, method)()
                except Exception:
                    value = None
                rows.append((title, value or "n/a"))
            try:
                conn.read_sizes()
                rows += [("Users", conn.users), ("Fingerprints", conn.fingers),
                         ("Faces", getattr(conn, "faces", "n/a")), ("Attendance records", conn.records)]
            except Exception:
                pass
            try:
                dev_time, now = conn.get_time(), datetime.now()
                rows += [("Device time", dev_time.strftime("%Y-%m-%d %H:%M:%S")),
                         ("Server time", now.strftime("%Y-%m-%d %H:%M:%S"))]
                drift = (dev_time - now).total_seconds()
                if abs(drift) > 300:
                    note = (f"Device clock differs from this server by {_fmt_delta(drift)}. Expected if they use "
                            "different time zones; otherwise fix the device clock, since punch times come from it.")
            except Exception:
                pass
    except Exception as exc:
        return {"ok": False, "error": friendly_error(exc), "elapsed": round(time.monotonic() - started, 1)}
    return {"ok": True, "rows": rows, "note": note, "elapsed": round(time.monotonic() - started, 1)}


# --------------------------------------------------------------------------- output files

HEADERS = ["Device", "User ID", "Name", "Date", "Time", "Punch", "Verify Method"]


def _uid_key(uid):
    return (0, int(uid), "") if str(uid).isdigit() else (1, 0, str(uid))


def _style_sheet(ws, widths):
    fill = PatternFill("solid", start_color="1F2937")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center")
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions


def _append(ws, values):
    """Appends a row; text starting with = + - @ is forced to plain text (blocks formula injection)."""
    ws.append(values)
    for cell in ws[ws.max_row]:
        if isinstance(cell.value, str) and cell.value[:1] in ("=", "+", "-", "@"):
            cell.data_type = "s"


def daily_rows(records):
    groups = defaultdict(list)
    for r in records:
        groups[(r["user_id"], r["timestamp"].date())].append(r)
    rows = []
    for (uid, day), recs in groups.items():
        times = sorted(x["timestamp"] for x in recs)
        rows.append({
            "uid": uid, "name": next((x["name"] for x in recs if x["name"]), ""), "day": day,
            "first": times[0].time(), "last": times[-1].time(), "punches": len(times),
            "hours": round((times[-1] - times[0]).total_seconds() / 3600, 2),
            "devices": ", ".join(sorted({x["device"] for x in recs})),
        })
    rows.sort(key=lambda r: (r["day"], _uid_key(r["uid"])))
    return rows


def build_workbook(records, summary, start=None, end=None, source="live", shift_rows=None, shift_stats=None):
    wb = Workbook()

    ws = wb.active
    ws.title = "Attendance"
    ws.append(HEADERS)
    for r in records:
        ts = r["timestamp"]
        _append(ws, [r["device"], r["user_id"], r["name"], ts.date(), ts.time(), r["punch"], r["verify"]])
    for row in ws.iter_rows(min_row=2):
        row[3].number_format = "yyyy-mm-dd"
        row[4].number_format = "hh:mm:ss"
    _style_sheet(ws, [22, 12, 24, 12, 11, 14, 16])

    daily = daily_rows(records)
    ds = wb.create_sheet("Daily Summary")
    ds.append(["User ID", "Name", "Date", "First Punch", "Last Punch", "Punches", "Hours Between", "Devices"])
    for d in daily:
        _append(ds, [d["uid"], d["name"], d["day"], d["first"], d["last"], d["punches"], d["hours"], d["devices"]])
    for row in ds.iter_rows(min_row=2):
        row[2].number_format = "yyyy-mm-dd"
        row[3].number_format = "hh:mm:ss"
        row[4].number_format = "hh:mm:ss"
    _style_sheet(ds, [12, 24, 12, 12, 12, 10, 14, 30])

    totals = {}
    for d in daily:
        t = totals.setdefault(d["uid"], {"name": d["name"], "days": 0, "punches": 0, "hours": 0.0})
        t["name"] = t["name"] or d["name"]
        t["days"] += 1
        t["punches"] += d["punches"]
        t["hours"] += d["hours"]
    et = wb.create_sheet("Employee Totals")
    et.append(["User ID", "Name", "Days Present", "Total Punches", "Total Hours Between"])
    for uid in sorted(totals, key=_uid_key):
        t = totals[uid]
        _append(et, [uid, t["name"], t["days"], t["punches"], round(t["hours"], 2)])
    _style_sheet(et, [12, 24, 14, 14, 20])

    if shift_rows is not None:
        _add_shift_sheets(wb, shift_rows)

    dv = wb.create_sheet("Devices")
    dv.append(["Device", "Address", "In Range", "On Device", "Duplicates Removed", "Seconds", "Status / Error"])
    for s in summary:
        _append(dv, [s["name"], s["address"], s["count"], s.get("on_device", 0), s.get("duplicates", 0),
                     s.get("elapsed", 0), s["error"] or "OK"])
    _style_sheet(dv, [24, 28, 10, 10, 18, 9, 70])

    info = wb.create_sheet("Report Info")
    info.append(["Item", "Value"])
    for k, v in [("Generated", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
                 ("Date range", f"{start} to {end}" if start and end else "n/a"),
                 ("Source", "Local database" if source == "db" else "Live from devices"),
                 ("Total records", len(records)),
                 *([("Punches outside any shift window", (shift_stats or {}).get("unmatched_punches", 0)),
                    ("Users with no shift", ", ".join((shift_stats or {}).get("users_without_shift", [])) or "none")]
                   if shift_rows is not None else []),
                 ("Note", "Hours Between = last punch minus first punch of the day; breaks are not deducted. "
                          "Times are device local time.")]:
        _append(info, [k, v])
    _style_sheet(info, [16, 90])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _add_shift_sheets(wb, rows):
    ws = wb.create_sheet("Shift Report")
    ws.append(["Shift Date", "User ID", "Name", "Shift", "Scheduled In", "Scheduled Out", "Actual In", "Actual Out",
               "Late (min)", "Early Leave (min)", "Overtime (min)", "Hours Worked", "Status"])
    for r in rows:
        _append(ws, [r["date"], r["uid"], r["name"], r["shift"], r["sched_in"], r["sched_out"], r["in"], r["out"],
                     r["late_min"], r["early_min"], r["ot_min"], r["hours"], r["status"]])
    for row in ws.iter_rows(min_row=2):
        row[0].number_format = "yyyy-mm-dd"
        for i in (4, 5, 6, 7):
            row[i].number_format = "yyyy-mm-dd hh:mm"
    _style_sheet(ws, [12, 10, 22, 12, 17, 17, 17, 17, 10, 16, 14, 13, 24])

    st = wb.create_sheet("Shift Totals")
    st.append(["User ID", "Name", "Shifts Worked", "Absent", "Late Count", "Late (min)", "Early Leave Count",
               "Early Leave (min)", "Overtime (min)", "Hours Worked"])
    for t in shiftlib.totals(rows):
        _append(st, [t["uid"], t["name"], t["shifts"], t["absent"], t["late_count"], t["late_min"],
                     t["early_count"], t["early_min"], t["ot_min"], t["hours"]])
    _style_sheet(st, [10, 22, 13, 9, 11, 11, 17, 17, 14, 13])


def _csv_safe(value):
    s = str(value)
    return "'" + s if s and s[0] in "=+-@" else s


def build_csv(records):
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(HEADERS)
    for r in records:
        ts = r["timestamp"]
        writer.writerow([_csv_safe(r["device"]), _csv_safe(r["user_id"]), _csv_safe(r["name"]),
                         ts.strftime("%Y-%m-%d"), ts.strftime("%H:%M:%S"), r["punch"], r["verify"]])
    return ("\ufeff" + buf.getvalue()).encode("utf-8")  # BOM so Excel opens UTF-8 names correctly
