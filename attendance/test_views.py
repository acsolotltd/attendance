"""Django tests: python manage.py test   (demo device + simulated ADMS pushes, no hardware needed)."""
import datetime as dt
import io

import openpyxl
from django.contrib.auth import get_user_model
from django.core.cache import cache
from django.test import TestCase, override_settings
from django.urls import reverse

from .models import AttendanceLog, Device, DeviceCommand, DeviceRequestLog


class LoggedInTestCase(TestCase):
    def setUp(self):
        cache.clear()  # rate limits / login throttle live in the cache
        self.client.force_login(get_user_model().objects.create_user("tester", password="x"))


class ViewTests(LoggedInTestCase):
    def setUp(self):
        super().setUp()
        self.demo = Device.objects.create(name="Demo A", ip_address="demo")

    def payload(self, **extra):
        today = dt.date.today()
        data = {"devices": [self.demo.pk], "start_date": (today - dt.timedelta(days=14)).isoformat(),
                "end_date": today.isoformat(), "source": "live", "fmt": "xlsx", "save_copy": "on"}
        data.update(extra)
        return data

    def test_index_loads(self):
        self.assertContains(self.client.get(reverse("index")), "Demo A")

    def test_add_pull_and_adms_device_and_validation(self):
        base = {"name": "HQ", "mode": "pull", "ip_address": "1.2.3.4", "port": 4370, "comm_key": 0,
                "timeout": 15, "protocol": "tcp", "serial_number": ""}
        self.assertEqual(self.client.post(reverse("device_add"), base).status_code, 200)
        adms = dict(base, name="Push", mode="adms", ip_address="", serial_number="CJDN123456")
        self.client.post(reverse("device_add"), adms)
        self.assertEqual(Device.objects.count(), 3)
        dup = self.client.post(reverse("device_add"), adms)           # same serial again -> rejected
        self.assertEqual(dup["HX-Retarget"], "#add-device")
        bad = self.client.post(reverse("device_add"), dict(base, ip_address="1.2.3.4:4370"))
        self.assertEqual(bad["HX-Retarget"], "#add-device")
        self.assertEqual(Device.objects.count(), 3)

    def test_edit_and_delete(self):
        self.client.post(reverse("device_edit", args=[self.demo.pk]), {
            "name": "Renamed", "mode": "pull", "ip_address": "demo", "port": 4370, "comm_key": 0, "timeout": 15,
            "protocol": "tcp", "serial_number": ""})
        self.demo.refresh_from_db()
        self.assertEqual(self.demo.name, "Renamed")
        self.client.post(reverse("device_delete", args=[self.demo.pk]))
        self.assertEqual(Device.objects.count(), 0)

    def test_status_and_test_endpoints(self):
        self.assertContains(self.client.get(reverse("device_status", args=[self.demo.pk])), "Online")
        self.assertContains(self.client.get(reverse("device_test", args=[self.demo.pk])), "simulator")

    def test_preview_saves_copy(self):
        self.assertContains(self.client.post(reverse("preview"), self.payload()), "records")
        self.assertGreater(AttendanceLog.objects.count(), 0)

    def test_export_xlsx_and_csv(self):
        resp = self.client.post(reverse("export"), self.payload())
        self.assertIn("spreadsheetml", resp["Content-Type"])
        wb = openpyxl.load_workbook(io.BytesIO(resp.content))
        self.assertIn("Daily Summary", wb.sheetnames)
        self.assertIn("text/csv", self.client.post(reverse("export"), self.payload(fmt="csv"))["Content-Type"])

    def test_sync_is_idempotent_and_db_source_works(self):
        self.client.post(reverse("sync"), {"devices": [self.demo.pk]})
        first = AttendanceLog.objects.count()
        self.assertGreater(first, 0)
        self.client.post(reverse("sync"), {"devices": [self.demo.pk]})
        self.assertEqual(AttendanceLog.objects.count(), first)
        self.assertContains(self.client.post(reverse("preview"), self.payload(source="db")), "from local database")

    def test_export_when_every_device_fails_returns_502(self):
        dead = Device.objects.create(name="Dead", ip_address="127.0.0.1", port=1)
        self.assertEqual(self.client.post(reverse("export"), self.payload(devices=[dead.pk])).status_code, 502)


@override_settings(ADMS_AUTO_REGISTER=True)
class AdmsTests(LoggedInTestCase):
    SN = "TESTSN0001"

    def now_line(self, uid="101", minutes_ago=5, punch=0, verify=1):
        ts = (dt.datetime.now() - dt.timedelta(minutes=minutes_ago)).replace(microsecond=0)
        return f"{uid}\t{ts:%Y-%m-%d %H:%M:%S}\t{punch}\t{verify}\t0\t0\t0\t0\r\n"

    def push(self, body, **params):
        q = {"SN": self.SN, "table": "ATTLOG", **params}
        url = "/iclock/cdata?" + "&".join(f"{k}={v}" for k, v in q.items())
        return self.client.post(url, data=body, content_type="text/plain")

    def payload(self, device, **extra):
        today = dt.date.today()
        data = {"devices": [device.pk], "start_date": (today - dt.timedelta(days=1)).isoformat(),
                "end_date": today.isoformat(), "fmt": "xlsx"}
        data.update(extra)
        return data

    def test_handshake_auto_registers_and_answers_options(self):
        resp = self.client.get("/iclock/cdata", {"SN": self.SN, "options": "all", "pushver": "2.4.1"})
        self.assertEqual(resp.status_code, 200)
        self.assertTrue(resp.content.decode().startswith(f"GET OPTION FROM: {self.SN}"))
        device = Device.objects.get(serial_number=self.SN)
        self.assertEqual(device.mode, "adms")
        self.assertIsNotNone(device.last_seen)

    @override_settings(ADMS_AUTO_REGISTER=False)
    def test_unknown_device_rejected_when_auto_register_off(self):
        self.assertEqual(self.client.get("/iclock/cdata", {"SN": self.SN}).status_code, 403)
        self.assertEqual(Device.objects.count(), 0)

    def test_invalid_serial_rejected(self):
        self.assertEqual(self.client.get("/iclock/cdata", {"SN": "../x"}).status_code, 400)

    def test_attlog_is_stored_idempotent_and_stamp_returned(self):
        body = self.now_line("101") + self.now_line("102", 3, punch=1)
        self.assertEqual(self.push(body, Stamp="777").content, b"OK")
        self.assertEqual(AttendanceLog.objects.count(), 2)
        self.push(body, Stamp="777")                                   # device re-sends: no duplicates
        self.assertEqual(AttendanceLog.objects.count(), 2)
        handshake = self.client.get("/iclock/cdata", {"SN": self.SN, "options": "all"}).content.decode()
        self.assertIn("ATTLOGStamp=777", handshake)

    def test_unparseable_attlog_is_not_acknowledged(self):
        self.assertEqual(self.push("this is not a punch\r\n").status_code, 400)
        self.assertEqual(AttendanceLog.objects.count(), 0)

    def test_user_names_fill_existing_and_new_punches(self):
        self.push(self.now_line("101"))
        users = "USER PIN=101\tName=Sade Williams\tPri=0\r\n"
        self.push(users, table="OPERLOG")
        self.assertEqual(AttendanceLog.objects.get().name, "Sade Williams")
        self.push(self.now_line("101", 1))
        self.assertEqual(set(AttendanceLog.objects.values_list("name", flat=True)), {"Sade Williams"})

    def test_command_round_trip(self):
        self.client.get("/iclock/cdata", {"SN": self.SN, "options": "all"})
        device = Device.objects.get(serial_number=self.SN)
        self.assertEqual(self.client.get("/iclock/getrequest", {"SN": self.SN}).content, b"OK")
        self.client.post(reverse("device_command", args=[device.pk]), {"cmd": "info"})
        cmd = DeviceCommand.objects.get()
        polled = self.client.get("/iclock/getrequest", {"SN": self.SN}).content.decode()
        self.assertEqual(polled, f"C:{cmd.pk}:INFO\n")
        self.assertEqual(self.client.get("/iclock/getrequest", {"SN": self.SN}).content, b"OK")  # delivered once
        self.client.post(f"/iclock/devicecmd?SN={self.SN}", data=f"ID={cmd.pk}&Return=0&CMD=INFO",
                         content_type="text/plain")
        cmd.refresh_from_db()
        self.assertEqual(cmd.result_code, 0)
        self.assertIsNotNone(cmd.finished_at)

    def test_unsafe_commands_cannot_be_queued_from_ui(self):
        self.client.get("/iclock/cdata", {"SN": self.SN})
        device = Device.objects.get(serial_number=self.SN)
        resp = self.client.post(reverse("device_command", args=[device.pk]), {"cmd": "DATA DELETE USERINFO PIN=1"})
        self.assertEqual(resp.status_code, 400)
        self.assertEqual(DeviceCommand.objects.count(), 0)

    def test_status_details_resync_and_export(self):
        self.push(self.now_line("101"), Stamp="55")
        device = Device.objects.get(serial_number=self.SN)
        self.assertContains(self.client.get(reverse("device_status", args=[device.pk])), "Online")
        self.assertContains(self.client.get(reverse("device_test", args=[device.pk])), self.SN)
        resp = self.client.post(reverse("export"), self.payload(device))
        wb = openpyxl.load_workbook(io.BytesIO(resp.content))
        self.assertEqual(wb["Attendance"].max_row, 2)
        self.client.post(reverse("device_resync", args=[device.pk]))
        device.refresh_from_db()
        self.assertEqual(device.att_stamp, "None")

    def test_requests_are_journaled_and_catch_all_says_ok(self):
        self.assertEqual(self.client.get("/iclock/ping").content, b"OK")
        self.client.get("/iclock/cdata", {"SN": self.SN})
        self.assertGreaterEqual(DeviceRequestLog.objects.count(), 2)
        self.assertContains(self.client.get(reverse("adms_log")), "/iclock/cdata")

    def test_ui_requires_login_but_devices_do_not(self):
        self.client.logout()
        self.assertEqual(self.client.get(reverse("index")).status_code, 302)
        self.assertEqual(self.client.get(reverse("index"), HTTP_HX_REQUEST="true").status_code, 401)
        self.assertEqual(self.client.get("/iclock/cdata", {"SN": self.SN}).status_code, 200)


@override_settings(ADMS_AUTO_REGISTER=True)
class ShiftViewTests(LoggedInTestCase):
    SN = "SHIFTSN001"

    def setUp(self):
        super().setUp()
        self.day0 = dt.date.today() - dt.timedelta(days=3)   # the night shift STARTS on this date

    def push_night_punches(self):
        """Pushes 19:58 on day0 and 06:07 the next morning, i.e. one night shift across midnight."""
        lines = (f"101\t{self.day0} 19:58:00\t0\t1\t0\r\n"
                 f"101\t{self.day0 + dt.timedelta(days=1)} 06:07:00\t1\t1\t0\r\n")
        self.client.post(f"/iclock/cdata?SN={self.SN}&table=ATTLOG", data=lines, content_type="text/plain")
        return Device.objects.get(serial_number=self.SN)

    def test_seed_add_edit_assign_delete(self):
        from .models import Shift, ShiftAssignment
        self.assertContains(self.client.get(reverse("shifts_panel")), "Create example")
        self.client.post(reverse("shift_seed"))
        self.assertEqual(set(Shift.objects.values_list("name", flat=True)), {"Day", "Night"})
        night = Shift.objects.get(name="Night")
        self.assertTrue(night.crosses_midnight)
        bad = self.client.post(reverse("shift_add"), {"name": "Zero", "start_time": "08:00", "end_time": "08:00",
                                                       "work_days": ["0"], "grace_minutes": 10,
                                                       "overtime_after_minutes": 30, "window_before_minutes": 120,
                                                       "window_after_minutes": 180})
        self.assertContains(bad, "can&#x27;t be the same")
        self.client.post(reverse("shift_edit", args=[night.pk]), {
            "name": "Night", "start_time": "21:00", "end_time": "06:00", "work_days": ["0", "1"], "grace_minutes": 5,
            "overtime_after_minutes": 30, "window_before_minutes": 120, "window_after_minutes": 180})
        night.refresh_from_db()
        self.assertEqual((night.start_time.hour, night.work_days), (21, "0,1"))
        self.client.post(reverse("assign_add"), {"shift": night.pk, "user_ids": "101, 102", "start_date": str(self.day0),
                                                  "rotate_with": "", "rotate_weeks": 1})
        self.assertEqual(ShiftAssignment.objects.count(), 2)
        self.client.post(reverse("assign_delete", args=[ShiftAssignment.objects.first().pk]))
        self.assertEqual(ShiftAssignment.objects.count(), 1)
        self.client.post(reverse("shift_delete", args=[night.pk]))
        self.assertEqual(ShiftAssignment.objects.count(), 0)

    def test_rotation_creates_alternating_blocks(self):
        from .models import Shift, ShiftAssignment
        self.client.post(reverse("shift_seed"))
        day, night = Shift.objects.get(name="Day"), Shift.objects.get(name="Night")
        self.client.post(reverse("assign_add"), {
            "shift": day.pk, "user_ids": "101", "start_date": str(self.day0),
            "end_date": str(self.day0 + dt.timedelta(days=20)), "rotate_with": night.pk, "rotate_weeks": 1})
        shifts_in_order = [a.shift.name for a in ShiftAssignment.objects.order_by("id")]
        self.assertEqual(shifts_in_order, ["Day", "Night", "Day"])

    def test_export_has_night_shift_row_spanning_midnight(self):
        from .models import Shift, ShiftAssignment
        device = self.push_night_punches()
        self.client.post(reverse("shift_seed"))
        night = Shift.objects.get(name="Night")
        ShiftAssignment.objects.create(user_id="101", shift=night, start_date=self.day0 - dt.timedelta(days=7))
        payload = {"devices": [device.pk], "start_date": str(self.day0), "end_date": str(self.day0),
                   "fmt": "xlsx", "shift_report": "on"}
        resp = self.client.post(reverse("export"), payload)
        wb = openpyxl.load_workbook(io.BytesIO(resp.content))
        self.assertIn("Shift Report", wb.sheetnames)
        rows = [r for r in wb["Shift Report"].iter_rows(min_row=2, values_only=True) if r[12] != "Absent"]
        self.assertEqual(len(rows), 1)
        self.assertEqual((rows[0][1], rows[0][3], rows[0][12]), ("101", "Night", "On time"))
        self.assertEqual(rows[0][7].date(), self.day0 + dt.timedelta(days=1))       # check-out is the next calendar day
        self.assertEqual(wb["Attendance"].max_row, 2)                                 # only the 19:58 punch is in range
        preview = self.client.post(reverse("preview"), payload)
        self.assertContains(preview, "Shift report")



class ProductionBehaviourTests(LoggedInTestCase):
    SN = "PRODSN0001"

    def test_unknown_device_is_listed_then_registered_with_one_click(self):
        self.assertEqual(self.client.get("/iclock/cdata", {"SN": self.SN}).status_code, 403)   # default: refuse
        self.assertContains(self.client.get(reverse("pending")), self.SN)
        self.client.post(reverse("device_register"), {"serial": self.SN})
        self.assertEqual(self.client.get("/iclock/cdata", {"SN": self.SN}).status_code, 200)
        self.assertNotContains(self.client.get(reverse("pending")), self.SN)

    @override_settings(ADMS_AUTO_REGISTER=True)
    def test_binary_uploads_are_acknowledged_but_not_stored(self):
        resp = self.client.post(f"/iclock/cdata?SN={self.SN}&table=ATTPHOTO&Stamp=1", data=b"\xff\xd8binary",
                                content_type="application/octet-stream")
        self.assertEqual(resp.content, b"OK")
        log = DeviceRequestLog.objects.filter(path__contains="ATTPHOTO").first()
        self.assertIn("not stored", log.body_preview)
        self.assertEqual(AttendanceLog.objects.count(), 0)

    @override_settings(ADMS_AUTO_REGISTER=True)
    def test_heartbeats_are_not_journaled_but_errors_are(self):
        self.client.get("/iclock/cdata", {"SN": self.SN})
        before = DeviceRequestLog.objects.count()
        self.client.get("/iclock/getrequest", {"SN": self.SN})
        self.assertEqual(DeviceRequestLog.objects.count(), before)

    @override_settings(ADMS_AUTO_REGISTER=True)
    def test_gbk_names_and_push3_flow(self):
        self.client.get("/iclock/cdata", {"SN": self.SN})
        body = "USER PIN=101\tName=张伟\tPri=0\r\n".encode("gb18030")
        self.client.post(f"/iclock/cdata?SN={self.SN}&table=OPERLOG", data=body, content_type="text/plain")
        from .models import DeviceUser
        self.assertEqual(DeviceUser.objects.get().name, "张伟")
        reg = self.client.post(f"/iclock/registry?SN={self.SN}", data="~DeviceName=MB360 Plus,StringPinFunOn=1",
                               content_type="text/plain")
        self.assertTrue(reg.content.decode().startswith("RegistryCode="))
        self.assertIn("SessionID=", self.client.get("/iclock/push", {"SN": self.SN}).content.decode())
        rows = f"transaction pin=101,verified=15,inoutstate=0,time={dt.datetime.now():%Y-%m-%d %H:%M:%S}\n"
        self.client.post(f"/iclock/cdata?SN={self.SN}&table=tabledata&tablename=transaction", data=rows,
                         content_type="text/plain")
        self.assertEqual(AttendanceLog.objects.get().verify, 15)

    @override_settings(ADMS_AUTO_REGISTER=True)
    def test_upload_profile_changes_handshake(self):
        self.client.get("/iclock/cdata", {"SN": self.SN})
        device = Device.objects.get(serial_number=self.SN)
        device.trans_profile = "minimal"
        device.save()
        text = self.client.get("/iclock/cdata", {"SN": self.SN}).content.decode()
        self.assertIn("TransFlag=TransData AttLog OpLog EnrollUser ChgUser", text)

    @override_settings(ADMS_AUTO_REGISTER=True, ADMS_ALLOWED_NETWORKS=[__import__("ipaddress").ip_network("10.9.9.0/24")])
    def test_ip_allow_list_blocks_other_sources(self):
        self.assertEqual(self.client.get("/iclock/cdata", {"SN": self.SN}, REMOTE_ADDR="198.51.100.5").status_code, 403)
        self.assertEqual(self.client.get("/iclock/cdata", {"SN": self.SN}, REMOTE_ADDR="10.9.9.4").status_code, 200)

    @override_settings(ADMS_AUTO_REGISTER=True, ADMS_RATE_LIMIT_PER_MIN=3)
    def test_rate_limit(self):
        codes = [self.client.get("/iclock/cdata", {"SN": self.SN}).status_code for _ in range(5)]
        self.assertEqual(codes[:3], [200, 200, 200])
        self.assertEqual(codes[-1], 429)

    def test_dashboard_health_and_stats(self):
        self.assertEqual(self.client.get("/healthz").json(), {"status": "ok"})
        self.assertContains(self.client.get(reverse("stats")), "Punches today")
        self.client.logout()
        self.assertEqual(self.client.get("/healthz").status_code, 200)   # health check needs no login

    def test_login_is_throttled_after_repeated_failures(self):
        self.client.logout()
        get_user_model().objects.create_user("bob", password="correct-horse-battery")
        for _ in range(8):
            self.client.post("/accounts/login/", {"username": "bob", "password": "nope"})
        self.assertEqual(self.client.post("/accounts/login/", {"username": "bob", "password": "correct-horse-battery"}).status_code, 429)
