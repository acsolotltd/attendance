"""Django-free tests: run with  python -m unittest attendance.test_services  (needs only openpyxl)."""
import datetime as dt
import io
import sys
import types
import unittest
from types import SimpleNamespace
from unittest import mock

import openpyxl

from attendance import services


def dev(**kw):
    base = dict(pk=1, name="HQ", ip_address="demo", port=4370, comm_key=0, timeout=5, protocol="tcp")
    base.update(kw)
    return SimpleNamespace(**base)


def log(uid, ts, punch=0, status=1):
    return SimpleNamespace(user_id=uid, timestamp=ts, punch=punch, status=status)


def fake_zk(attendance, users=None, fail_connect=0, exc=None):
    state = {"connects": 0, "enabled": True, "disconnected": False}

    class FakeZK:
        def __init__(self, host, port=4370, timeout=60, password=0, force_udp=False, ommit_ping=False):
            self.kw = dict(host=host, port=port, timeout=timeout, password=password, force_udp=force_udp)
            self.users, self.records, self.fingers, self.faces = 2, len(attendance), 3, 0

        def connect(self):
            state["connects"] += 1
            if exc and state["connects"] <= fail_connect:
                raise exc
            return self

        def disable_device(self): state["enabled"] = False
        def enable_device(self): state["enabled"] = True
        def disconnect(self): state["disconnected"] = True
        def get_users(self): return [SimpleNamespace(user_id=k, name=v) for k, v in (users or {}).items()]
        def get_attendance(self): return list(attendance)
        def get_time(self): return dt.datetime.now()
        def get_serialnumber(self): return "SN123"
        def get_firmware_version(self): return "Ver 6.60"
        def get_platform(self): return "ZMM220"
        def get_device_name(self): return "MB360"
        def get_mac(self): return "00:11:22:33:44:55"
        def read_sizes(self): pass

    mod = types.ModuleType("zk")
    mod.ZK = FakeZK
    return mod, state


class DemoTests(unittest.TestCase):
    def test_demo_range_and_duplicates(self):
        today = dt.date.today()
        recs, meta = services.fetch_device_records(dev(), today - dt.timedelta(days=30), today)
        self.assertIsNone(meta["error"])
        self.assertTrue(recs)
        self.assertGreaterEqual(meta["duplicates"], 1)
        self.assertTrue(all(today - dt.timedelta(days=30) <= r["timestamp"].date() <= today for r in recs))
        self.assertEqual(recs, sorted(recs, key=lambda r: (r["timestamp"], r["user_id"])))

    def test_user_filter_by_id_and_name(self):
        today = dt.date.today()
        by_id, _ = services.fetch_device_records(dev(), today - dt.timedelta(days=30), today, ["101"])
        self.assertTrue(by_id and {r["user_id"] for r in by_id} == {"101"})
        by_name, _ = services.fetch_device_records(dev(), today - dt.timedelta(days=30), today, ["chidi"])
        self.assertTrue(by_name and {r["name"] for r in by_name} == {"Chidi Okafor"})

    def test_parse_user_filter(self):
        self.assertEqual(services.parse_user_filter("101, 105;John Smith\n"), ["101", "105", "john smith"])


class DeviceTests(unittest.TestCase):
    def run_fetch(self, attendance, **kw):
        mod, state = fake_zk(attendance, users={"1": "Ann"}, **kw)
        with mock.patch.dict(sys.modules, {"zk": mod}), \
             mock.patch.object(services, "probe", return_value=(True, 5)), \
             mock.patch.object(services.time, "sleep"):
            out = services.fetch_device_records(dev(ip_address="1.2.3.4"), dt.date(2026, 10, 1), dt.date(2026, 10, 2),
                                                keep_all=True)
        return out, state

    def test_filters_dedupes_and_restores_device(self):
        logs = [log("1", dt.datetime(2026, 10, 1, 8, 0, 0)), log("1", dt.datetime(2026, 10, 1, 8, 0, 0)),
                log("1", dt.datetime(2026, 10, 1, 17, 0, 0), punch=1), log("1", dt.datetime(2026, 9, 1, 8, 0, 0))]
        (recs, meta), state = self.run_fetch(logs)
        self.assertIsNone(meta["error"])
        self.assertEqual((len(recs), meta["on_device"], meta["duplicates"]), (2, 3, 1))
        self.assertEqual(len(meta["all_records"]), 3)
        self.assertEqual(recs[0]["name"], "Ann")
        self.assertEqual(recs[1]["punch"], "Check Out")
        self.assertTrue(state["enabled"] and state["disconnected"])

    def test_retry_then_success(self):
        (recs, meta), state = self.run_fetch([log("1", dt.datetime(2026, 10, 1, 8, 0))],
                                             fail_connect=1, exc=OSError("boom"))
        self.assertIsNone(meta["error"])
        self.assertEqual(state["connects"], 2)

    def test_auth_failure_is_not_retried_and_hints_comm_key(self):
        (recs, meta), state = self.run_fetch([], fail_connect=5, exc=Exception("Unauthenticated"))
        self.assertEqual(state["connects"], 1)
        self.assertIn("Comm Key", meta["error"])

    def test_unreachable_tcp(self):
        with mock.patch.object(services, "probe", return_value=(False, None)):
            recs, meta = services.fetch_device_records(dev(ip_address="9.9.9.9"), dt.date(2026, 10, 1), dt.date(2026, 10, 2))
        self.assertEqual(recs, [])
        self.assertIn("Cannot reach", meta["error"])

    def test_punch_state_unused_labels_everything_punch(self):
        mod, _ = fake_zk([log("1", dt.datetime(2026, 10, 1, 8, 0), punch=4)], users={"1": "Ann"})
        with mock.patch.dict(sys.modules, {"zk": mod}), \
             mock.patch.object(services, "probe", return_value=(True, 5)), mock.patch.object(services.time, "sleep"):
            recs, _ = services.fetch_device_records(dev(ip_address="1.2.3.4", punch_state_unused=True),
                                                    dt.date(2026, 10, 1), dt.date(2026, 10, 1))
            plain, _ = services.fetch_device_records(dev(ip_address="1.2.3.4"), dt.date(2026, 10, 1), dt.date(2026, 10, 1))
        self.assertEqual((recs[0]["punch"], recs[0]["punch_code"]), ("Punch", 4))
        self.assertEqual(plain[0]["punch"], "Overtime In")

    def test_old_pyzk_without_encoding_argument_still_works(self):
        mod, state = fake_zk([log("1", dt.datetime(2026, 10, 1, 8, 0))])  # FakeZK rejects encoding= like old pyzk
        with mock.patch.dict(sys.modules, {"zk": mod}):
            zk = services._make_zk(dev(ip_address="1.2.3.4", encoding="GBK"))
        self.assertEqual(zk.kw["host"], "1.2.3.4")

    def test_device_info(self):
        mod, _ = fake_zk([])
        with mock.patch.dict(sys.modules, {"zk": mod}), mock.patch.object(services, "probe", return_value=(True, 5)):
            info = services.device_info(dev(ip_address="1.2.3.4"))
        self.assertTrue(info["ok"])
        self.assertIn(("Serial number", "SN123"), info["rows"])


class OutputTests(unittest.TestCase):
    def rec(self, device, uid, name, ts):
        return {"device": device, "user_id": uid, "name": name, "timestamp": ts, "punch": "Check In", "punch_code": 0,
                "verify": "Face", "verify_code": 15}

    def test_workbook_sheets_summary_and_formula_safety(self):
        recs = [self.rec("HQ", "1", "=HYPERLINK(\"x\")", dt.datetime(2026, 10, 1, 8, 0)),
                self.rec("Gate", "1", "=HYPERLINK(\"x\")", dt.datetime(2026, 10, 1, 17, 30))]
        buf = services.build_workbook(recs, [{"name": "HQ", "address": "a", "count": 1, "error": None}],
                                      dt.date(2026, 10, 1), dt.date(2026, 10, 1))
        wb = openpyxl.load_workbook(io.BytesIO(buf.getvalue()))
        self.assertEqual(wb.sheetnames, ["Attendance", "Daily Summary", "Employee Totals", "Devices", "Report Info"])
        self.assertEqual(wb["Attendance"]["C2"].data_type, "s")  # text, not a formula
        ds = list(wb["Daily Summary"].iter_rows(min_row=2, values_only=True))
        self.assertEqual(len(ds), 1)  # same user, same day, two devices -> one row
        self.assertEqual((ds[0][5], ds[0][6], ds[0][7]), (2, 9.5, "Gate, HQ"))

    def test_csv(self):
        out = services.build_csv([self.rec("HQ", "1", "-cmd", dt.datetime(2026, 10, 1, 8, 0))])
        self.assertTrue(out.startswith("\ufeff".encode("utf-8")))
        self.assertIn(b"'-cmd", out)


if __name__ == "__main__":
    unittest.main()
